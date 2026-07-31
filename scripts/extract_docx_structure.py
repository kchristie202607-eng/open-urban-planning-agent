from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document

parser = argparse.ArgumentParser()
parser.add_argument("input")
parser.add_argument("output")
args = parser.parse_args()

doc = Document(args.input)
lines: list[str] = []
for idx, paragraph in enumerate(doc.paragraphs, start=1):
    text = paragraph.text.strip()
    if text:
        lines.append(f"P{idx:04d}\t[{paragraph.style.name}]\t{text}")
for table_idx, table in enumerate(doc.tables, start=1):
    lines.append(f"\n===== TABLE {table_idx} ({len(table.rows)}x{len(table.columns)}) =====")
    for row in table.rows:
        lines.append("\t".join(cell.text.replace("\n", " | ").strip() for cell in row.cells))

Path(args.output).write_text("\n".join(lines), encoding="utf-8")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} output={args.output}")
