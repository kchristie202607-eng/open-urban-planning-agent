from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "output" / "demo"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "203748"
LIGHT_GRAY = "F2F4F7"
BLUE_GRAY = "E8EEF5"
MUTED = "666666"
GOLD = "8A6D1D"


def set_run_font(
    run, name: str = "Calibri", size: float | None = None,
    color: str | None = None, bold: bool | None = None,
    italic: bool | None = None
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "等线")
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def format_table(table, widths_dxa: list[int], header: bool = True) -> None:
    set_table_geometry(table, widths_dxa)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            if header and row_index == 0:
                set_cell_shading(cell, BLUE_GRAY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=9.2,
                        bold=(header and row_index == 0),
                        color=NAVY,
                    )


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run_font(run, size=9, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    list_style = styles["List Bullet"]
    list_style.paragraph_format.left_indent = Inches(0.5)
    list_style.paragraph_format.first_line_indent = Inches(-0.25)
    list_style.paragraph_format.space_after = Pt(8)
    list_style.paragraph_format.line_spacing = 1.167

    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    run = header_p.add_run("城市更新片区策划 | 本地知识库示范成果")
    set_run_font(run, size=9, color=MUTED)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_p.add_run("第 ")
    add_page_field(footer_p)
    footer_p.add_run(" 页")
    for run in footer_p.runs:
        set_run_font(run, size=9, color=MUTED)


def add_cover(doc: Document, project: dict, quality: dict) -> None:
    for _ in range(4):
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(14)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("片区策划完整工作流 - 示例成果")
    set_run_font(run, size=10.5, color=GOLD, bold=True)
    kicker.paragraph_format.space_after = Pt(18)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run(project["project_name"])
    set_run_font(run, size=30, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("片区更新策划与指标测算报告")
    set_run_font(run, size=15, color=DARK_BLUE)

    meta = doc.add_table(rows=4, cols=2)
    values = [
        ("项目编号", project["project_id"]),
        ("区位", project["location"]),
        ("成果状态", quality["status"]),
        ("成果属性", "策划阶段示范，不替代法定规划和专项审批"),
    ]
    for row, values_row in zip(meta.rows, values):
        row.cells[0].text = values_row[0]
        row.cells[1].text = values_row[1]
    format_table(meta, [2700, 6660], header=False)
    for row in meta.rows:
        set_cell_shading(row.cells[0], LIGHT_GRAY)
        row.cells[0].paragraphs[0].runs[0].bold = True

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("基于两份本地 PDF 知识库、可追溯检索与参数化测算生成")
    set_run_font(run, size=9.5, color=MUTED, italic=True)
    doc.add_page_break()


def add_callout(doc: Document, title: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = ""
    set_table_geometry(table, [9360])
    set_cell_shading(table.cell(0, 0), LIGHT_GRAY)
    paragraph = table.cell(0, 0).paragraphs[0]
    run = paragraph.add_run(title + "  ")
    set_run_font(run, size=10.5, color=DARK_BLUE, bold=True)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_metrics_table(doc: Document, metrics: dict) -> None:
    rows = [
        ("总用地面积", metrics["site_area_m2"], "m²"),
        ("目标容积率 / 上限", f"{metrics['target_far']:.2f} / {metrics['max_far']:.2f}", "-"),
        ("总建筑面积", metrics["total_gfa_m2"], "m²"),
        ("新增建筑面积", metrics["new_gfa_m2"], "m²"),
        ("建筑密度 / 上限", f"{metrics['building_density']:.1%} / {metrics['max_building_density']:.1%}", "%"),
        ("平均层数 / 高度", f"{metrics['average_storeys']} / {metrics['average_height_m']:.1f}", "层 / m"),
        ("绿地面积 / 绿地率", f"{metrics['green_area_m2']:,.0f} / {metrics['green_ratio']:.1%}", "m² / %"),
        ("公共空间", metrics["public_space_m2"], "m²"),
        ("停车位", metrics["parking_spaces"], "个"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "指标"
    table.rows[0].cells[1].text = "测算值"
    table.rows[0].cells[2].text = "单位"
    for label, value, unit in rows:
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = f"{value:,.0f}" if isinstance(value, (int, float)) else str(value)
        cells[2].text = unit
    format_table(table, [4200, 3300, 1860])


def build(output_dir: Path = DEFAULT_OUTPUT) -> Path:
    scenario = json.loads((output_dir / "scenario.json").read_text(encoding="utf-8"))
    evidence = json.loads((output_dir / "evidence.json").read_text(encoding="utf-8"))
    quality = json.loads((output_dir / "quality_report.json").read_text(encoding="utf-8"))
    with (output_dir / "project_list.csv").open(encoding="utf-8-sig", newline="") as handle:
        projects = list(csv.DictReader(handle))

    project = scenario["project"]
    metrics = scenario["metrics"]
    finance = scenario["finance"]
    assumptions = scenario["assumption_register"]

    doc = Document()
    configure_document(doc)
    add_cover(doc, project, quality)

    doc.add_heading("1. 结论摘要", level=1)
    add_callout(
        doc,
        "建议结论",
        (
            f"按当前策划假设，片区目标容积率为 {metrics['target_far']:.2f}，"
            f"总建筑面积约 {metrics['total_gfa_m2']:,.0f} 平方米，"
            f"总投资约 {finance['total_investment_yuan'] / 1e8:.2f} 亿元。"
            "控制指标和成本收益参数尚需法定依据与专项论证确认。"
        ),
    )
    p = doc.add_paragraph()
    p.add_run("策划主线：").bold = True
    p.add_run("、".join(project["planning_intents"]) + "。")
    p = doc.add_paragraph()
    p.add_run("成果边界：").bold = True
    p.add_run(
        "本报告用于片区策划阶段方案比较，不替代城市体检、控规、建筑方案、"
        "测绘、评估、专项审批或投资决策。"
    )

    doc.add_heading("2. 知识库与工作方法", level=1)
    for text in (
        "库 A 汇集《城市更新规划编制工作手册》，以编制程序、片区体检、策划内容和成果要求为主。",
        "库 B 汇集贵州省城市更新行动规划征求意见稿，以地方任务、项目类型、实施机制和目标指标为主。",
        "每条检索结果保留文件名、PDF 页码、抽取方式、文件状态和原文，模型不得将候选证据自动认定为强制要求。",
        "工作流依次执行证据检索、约束审查、空间指标测算、财务与项目库测算、成果质检。",
    ):
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("3. 片区策划策略", level=1)
    strategy_rows = [
        ("存量住区", "安全韧性、公共服务和环境品质综合提升", "近期启动"),
        ("完整社区", "补齐社区服务、公共空间与停车设施", "近期启动"),
        ("历史文化", "资源普查、分类保护、活化利用与运营导入", "重点实施"),
        ("产业空间", "低效空间盘活、功能置换和就业导入", "重点实施"),
    ]
    table = doc.add_table(rows=1, cols=3)
    for index, value in enumerate(("策略方向", "主要行动", "实施时序")):
        table.rows[0].cells[index].text = value
    for item in strategy_rows:
        cells = table.add_row().cells
        for index, value in enumerate(item):
            cells[index].text = value
    format_table(table, [2200, 5160, 2000])

    doc.add_heading("4. 指标测算", level=1)
    add_metrics_table(doc, metrics)
    for note in metrics["calculation_notes"]:
        doc.add_paragraph(note, style="List Bullet")

    doc.add_heading("5. 功能业态", level=1)
    table = doc.add_table(rows=1, cols=3)
    for index, value in enumerate(("功能", "配比", "建筑面积（m²）")):
        table.rows[0].cells[index].text = value
    for name, value in metrics["program_gfa_m2"].items():
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = f"{assumptions['program_mix'][name]:.1%}"
        cells[2].text = f"{value:,.0f}"
    format_table(table, [3900, 2000, 3460])

    doc.add_heading("6. 投资与资金平衡", level=1)
    finance_rows = [
        ("新建工程", finance["new_construction_cost_yuan"]),
        ("存量改造", finance["renovation_cost_yuan"]),
        ("公共空间", finance["public_space_cost_yuan"]),
        ("其他费用", finance["other_cost_yuan"]),
        ("预备费", finance["contingency_yuan"]),
        ("总投资", finance["total_investment_yuan"]),
        ("直接收益", finance["direct_revenue_yuan"]),
        ("资金缺口", finance["funding_gap_yuan"]),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "金额（亿元）"
    for label, value in finance_rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = f"{value / 1e8:.2f}"
    format_table(table, [5700, 3660])
    add_callout(doc, "测算提示", finance["warning"])

    doc.add_heading("7. 项目库", level=1)
    table = doc.add_table(rows=1, cols=5)
    headers = ("编码", "项目名称", "策划目标", "时序", "更新方式")
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for item in projects:
        cells = table.add_row().cells
        values = (
            item["project_code"],
            item["project_name"],
            item["planning_intent"],
            item["implementation_stage"],
            item["renewal_method"],
        )
        for index, value in enumerate(values):
            cells[index].text = value
    format_table(table, [1650, 2450, 2200, 1500, 1560])

    doc.add_heading("8. 证据索引", level=1)
    intro = doc.add_paragraph()
    intro.paragraph_format.space_before = Pt(4)
    intro.paragraph_format.space_after = Pt(4)
    run = intro.add_run("以下证据用于定位原文，正式引用前须核对对应 PDF 页面。")
    set_run_font(run, size=9.5, color=MUTED, italic=True)
    for index, item in enumerate(evidence[:10], start=1):
        p = doc.add_paragraph()
        p.style = doc.styles["List Bullet"]
        source = (
            f"[{index}] 《{item['source_file']}》PDF 第 {item['pdf_page']} 页；"
            f"{item['document_status']}；{item['extraction_method']}；"
            f"置信度 {item['confidence']:.3f}。"
        )
        run = p.add_run(source)
        set_run_font(run, size=9.5, color=DARK_BLUE, bold=True)
        quote = item["quotation"].replace("\n", " ")
        run = p.add_run(quote[:260] + ("…" if len(quote) > 260 else ""))
        set_run_font(run, size=9.5, color=NAVY)

    doc.add_heading("9. 假设、风险与下一步", level=1)
    for item in quality["review_items"]:
        doc.add_paragraph(item, style="List Bullet")

    output_dir.mkdir(parents=True, exist_ok=True)
    docx_path = output_dir / "片区更新策划与指标测算报告.docx"
    doc.save(docx_path)
    return docx_path


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    print(build(args.output.resolve()))
