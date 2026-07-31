from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import (
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_TABLE_ALIGNMENT,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUTPUT_DIR = Path("output/policy_framework")
OUTPUT_PATH = OUTPUT_DIR / "贵州省城市更新片区策划报告工作框架_V1.1_片区体检增强版.docx"

BLUE = "2E74B5"
NAVY = "1F4D78"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F4F8FB"
TEXT = "243240"
MUTED = "66717C"
RED = "C00000"
GREEN = "2F6B4F"
WHITE = "FFFFFF"
BORDER = "B7C4D0"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size=6) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_text(cell, text: str, *, bold=False, color=TEXT, size=9, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Calibri"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_field(paragraph, instruction: str) -> None:
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "打开文档后按 Ctrl+A、F9 更新目录"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([fld_char_begin, instr_text, fld_char_sep, text, fld_char_end])


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("— ")
    run.font.color.rgb = RGBColor.from_string(MUTED)
    add_field(paragraph, "PAGE")
    run = paragraph.add_run(" —")
    run.font.color.rgb = RGBColor.from_string(MUTED)


def set_run_font(run, east_asia="等线", latin="Calibri", size=None, color=None, bold=None) -> None:
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def set_outline_level(style, level: int) -> None:
    p_pr = style.element.get_or_add_pPr()
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        p_pr.append(outline)
    outline.set(qn("w:val"), str(level))


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_specs = {
        1: (16, BLUE, True, 18, 10),
        2: (13, BLUE, True, 14, 7),
        3: (12, NAVY, True, 10, 5),
        4: (11, NAVY, True, 8, 4),
        5: (10.5, MUTED, True, 6, 3),
        6: (10, MUTED, True, 4, 2),
    }
    for level, (size, color, bold, before, after) in heading_specs.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        set_outline_level(style, level - 1)

    for name, size, color, bold in (
        ("Framework Requirement", 9.5, TEXT, False),
        ("Evidence Line", 9, GREEN, False),
        ("Policy Line", 8.5, MUTED, False),
        ("Callout", 10, NAVY, False),
    ):
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.15
    doc.styles["Framework Requirement"].paragraph_format.left_indent = Cm(0.45)
    doc.styles["Framework Requirement"].paragraph_format.first_line_indent = Cm(-0.45)
    doc.styles["Evidence Line"].paragraph_format.left_indent = Cm(0.45)
    doc.styles["Evidence Line"].paragraph_format.first_line_indent = Cm(-0.45)
    doc.styles["Policy Line"].paragraph_format.left_indent = Cm(0.45)
    doc.styles["Policy Line"].paragraph_format.first_line_indent = Cm(-0.45)


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("贵州省城市更新片区策划｜报告工作框架")
    set_run_font(r, size=8.5, color=MUTED)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), BLUE)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    add_page_number(section.footer.paragraphs[0])


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(38)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("贵州省城市更新片区策划")
    set_run_font(r, east_asia="等线", size=15, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("报告工作框架")
    set_run_font(r, east_asia="等线", size=28, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("四级内容目录 · 六级标题样式预设 · 政策审查闭环")
    set_run_font(r, size=13, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("以“片区体检—策划响应—项目生成—资金平衡—运营实施”为主线，"
                  "用于策划方案编制、内审、联合评审及2027年省级财政资金支持项目申报。")
    set_run_font(r, size=11, color=TEXT)

    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = [Cm(3.3), Cm(12.2)]
    rows = [
        ("编制依据", "《贵州省城市更新片区策划方案审查要点（试行）》及2027年省级财政资金支持城市更新储备项目遴选通知"),
        ("适用范围", "贵州省城市更新重点片区策划方案；申报片区面积按通知不得低于10公顷"),
        ("框架深度", "正文实际展开至四级标题；Heading 1—Heading 6 样式已全部预设"),
        ("使用状态", "技术工作底稿 / Word可编辑模板 / 报审前逐项销项"),
        ("版本日期", "V1.1｜2026年7月｜片区体检增强版"),
    ]
    for row, (key, value) in zip(table.rows, rows):
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
        set_cell_text(row.cells[0], key, bold=True, color=NAVY, size=9)
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        set_cell_text(row.cells[1], value, size=9)
        prevent_row_split(row)
    set_table_borders(table)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("使用提示")
    set_run_font(r, size=10, color=BLUE, bold=True)
    p = doc.add_paragraph(style="Callout")
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.right_indent = Cm(0.35)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("本文件不是泛化提纲。每个四级标题均对应可核验的编制动作与成果证据；"
                  "“不得新增隐性债务、不得大拆大建、不得破坏历史文脉、不得以城市体检替代片区体检”"
                  "等事项应在提交评审前完成一票否决核验。")
    set_run_font(r, size=10, color=NAVY)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE_BLUE)
    p_pr.append(shd)

    doc.add_page_break()


def add_summary_tables(doc: Document) -> None:
    doc.add_heading("编制控制说明", level=1)
    p = doc.add_paragraph(
        "本框架以省厅《审查要点》附件1的七章参考大纲为刚性主目录，"
        "以附件2完整性、合规性及28项方案审查要求为验收清单，"
        "并叠加2027年试点通知的申报条件、竞争评审倾向、资金监管与时间节点。"
    )
    p.paragraph_format.space_after = Pt(8)

    doc.add_heading("一票否决与申报门槛", level=2)
    rows = [
        ("片区与主体", "边界清晰、完整连片；申报片区≥10公顷；明确统筹实施主体、运营主体并组建工作专班。", "范围批复/审核意见；主体文件；专班文件"),
        ("片区体检", "必须实质性开展片区体检，不得以既有城市体检替代；问题、资源、居民意愿、运营四类调查闭环。", "片区体检报告；数据库；入户台账；调查底图"),
        ("债务与资金", "不得新增地方政府隐性债务，不得随意扩大融资规模；整体资金平衡可验证。", "财政合规意见；资金平衡表；现金流与敏感性分析"),
        ("保护与建设", "不得大拆大建、破坏历史文脉或违规征拆；违法建筑调查后按需拆除释放公共空间。", "历史文化专项；违建台账；产权人意见；更新方式图"),
        ("规划合规", "符合国土空间总体规划和详细规划；拟调整详细规划须充分论证并征求部门意见。", "规划套合图；指标对照表；部门书面意见"),
        ("项目成熟度", "项目不得停留在意向阶段；需形成设计方案深度、投资估算、资金来源和分年度时序。", "项目一案一册；设计图；估算书；年度计划"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["控制主题", "明确工作要求", "最低验收证据"]
    widths = [Cm(2.6), Cm(8.2), Cm(5.0)]
    for i, h in enumerate(headers):
        table.columns[i].width = widths[i]
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=WHITE, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[i], BLUE)
    set_repeat_table_header(table.rows[0])
    for idx, values in enumerate(rows):
        row = table.add_row()
        for i, value in enumerate(values):
            row.cells[i].width = widths[i]
            set_cell_text(row.cells[i], value, bold=(i == 0), size=8.5)
            if idx % 2:
                set_cell_shading(row.cells[i], "F7F9FB")
        prevent_row_split(row)
    set_table_borders(table)

    doc.add_heading("标题层级与格式预设", level=2)
    p = doc.add_paragraph(
        "主报告仅使用一至四级标题；五、六级用于后续专项专篇、项目一案一册或复杂技术附录，"
        "样式已预设但不建议在总体报告中常态使用。"
    )
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["级别", "编号示例", "字号/颜色", "建议用途", "本框架"]
    widths = [Cm(1.3), Cm(3.2), Cm(3.2), Cm(6.1), Cm(2.0)]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=WHITE, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[i], NAVY)
    set_repeat_table_header(table.rows[0])
    style_rows = [
        ("一级", "第一章", "16pt / 蓝", "七章主目录", "使用"),
        ("二级", "1.1", "13pt / 蓝", "省厅参考大纲规定条目", "使用"),
        ("三级", "1.1.1", "12pt / 深蓝", "论证任务或专业模块", "使用"),
        ("四级", "1.1.1.1", "11pt / 深蓝", "可交付成果单元", "使用"),
        ("五级", "1.1.1.1.1", "10.5pt / 灰蓝", "专项分析或项目子任务", "预设"),
        ("六级", "1.1.1.1.1.1", "10pt / 灰蓝", "数据字段或技术细目", "预设"),
    ]
    for idx, values in enumerate(style_rows):
        row = table.add_row()
        for i, value in enumerate(values):
            set_cell_text(row.cells[i], value, bold=(i == 0), size=8.5)
            if idx % 2:
                set_cell_shading(row.cells[i], "F7F9FB")
        prevent_row_split(row)
    set_table_borders(table)

    doc.add_page_break()
    doc.add_heading("成果组织与审查闭环", level=2)
    flow = [
        ("片区体检", "形成问题、资源、意愿、运营四库及空间底图", "第二章"),
        ("策划响应", "每项问题形成目标、策略、空间、项目、资金、主体六项响应", "第三—七章"),
        ("项目生成", "项目达到设计方案深度，形成估算、资金来源、年度时序", "第五章"),
        ("合规内审", "完整性、合规性、28项审查要点逐项销项", "附录A—C"),
        ("申报实施", "支撑材料汇编、3个月开工、中期评估、年度绩效监测", "第七章及附录D"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(("阶段", "完成标准", "报告落位")):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=WHITE, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[i], BLUE)
    set_repeat_table_header(table.rows[0])
    for idx, values in enumerate(flow):
        row = table.add_row()
        for i, value in enumerate(values):
            set_cell_text(row.cells[i], value, bold=(i == 0), size=8.7)
            if idx % 2:
                set_cell_shading(row.cells[i], "F7F9FB")
        prevent_row_split(row)
    set_table_borders(table)

    doc.add_heading("城市体检方法引用", level=2)
    p = doc.add_paragraph(
        "本框架引用“2026-716-New project”中的《遵义片区体检项目任务书解析与初始化说明》"
        "及《遵义市2026年城市体检报告社区维度成果补充稿》作为工作方法支持。"
        "相关内容属于项目工作控制方法，不替代国家、省级政策和经确认的正式指标体系。"
    )
    method_lines = [
        ("四级衔接", "按住房—小区（社区）—街区—城区四级体检体系组织数据，片区策划重点承接小区（社区）和街区层级问题，并向具体空间、项目与运营场景转译。"),
        ("唯一事实源", "以指标底表、问题清单、更新项目库、数据来源与证据台账作为四张核心底表；报告、图件、汇报和项目表不得各自形成不同口径。"),
        ("证据分级", "A类政府正式统计/法定数据，B类部门业务/正规监测，C类GIS/遥感/现场调查，D类问卷/访谈/网络感知，E类AI或专家初判；E类不得直接写成确定性结论。"),
        ("数据准入", "任何资料进入计算前必须完成来源、年份、空间范围、统一主键和质量检查；原始文件只读留存，清洗结果不得覆盖原件。"),
        ("问题认定", "采用“指标异常＋空间异常＋群众诉求＋现场或部门证据”的四证合一原则；证据未闭合的标注为线索或待核实。"),
        ("阶段边界", "范围未确定时可开展资料盘点、数据治理、单元准备和候选片区筛选；不得形成正式指标结论、最终问题认定、完整居民调查结论或最终项目清单。"),
    ]
    for label, text in method_lines:
        p = doc.add_paragraph(style="Framework Requirement")
        r = p.add_run(label + "｜")
        set_run_font(r, size=9.5, color=BLUE, bold=True)
        r = p.add_run(text)
        set_run_font(r, size=9.5, color=TEXT)

    doc.add_heading("框架目录", level=1)
    p = doc.add_paragraph(
        "以下为省厅七章参考大纲及本模板二级目录。正文已采用真实 Heading 1—Heading 4 样式，"
        "可在 Word 导航窗格中展开至四级；正式成稿时可按需要插入自动目录。"
    )
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_cell_text(table.rows[0].cells[0], "一级目录", bold=True, color=WHITE, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.rows[0].cells[1], "二级目录", bold=True, color=WHITE, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(table.rows[0].cells[0], NAVY)
    set_cell_shading(table.rows[0].cells[1], NAVY)
    set_repeat_table_header(table.rows[0])
    for idx, (chapter_title, sections) in enumerate(CHAPTERS):
        row = table.add_row()
        section_names = "；".join(section_title for section_title, _ in sections)
        set_cell_text(row.cells[0], chapter_title, bold=True, color=NAVY, size=8.5)
        set_cell_text(row.cells[1], section_names, size=8.5)
        if idx % 2:
            set_cell_shading(row.cells[0], "F7F9FB")
            set_cell_shading(row.cells[1], "F7F9FB")
        prevent_row_split(row)
    set_table_borders(table)


def item(title: str, requirement: str, evidence: str, policy: str):
    return (title, requirement, evidence, policy)


CHAPTERS = [
    ("第一章 总则", [
        ("1.1 项目背景", [
            ("1.1.1 政策背景与更新必要性", [
                item("1.1.1.1 政策任务响应",
                     "逐条说明片区对省级城市更新、片区更新和财政支持政策的响应关系，区分强制要求、竞争评审事项和地方自选任务；不得仅罗列文件名称。",
                     "形成“政策条款—片区问题—策划动作—项目—责任部门”对照表。",
                     "《审查要点》第2、4—9页；《试点通知》第1—4页。"),
                item("1.1.1.2 片区更新必要性",
                     "用片区体检事实证明更新的必要性和紧迫性，至少覆盖民生短板、安全韧性、闲置低效、产业动能、历史文化与治理问题。",
                     "形成关键指标卡、问题热力图和不更新情景风险说明。",
                     "《审查要点》第3—4、7—8页。"),
            ]),
        ]),
        ("1.2 编制过程", [
            ("1.2.1 组织与技术路线", [
                item("1.2.1.1 工作组织",
                     "写明城市更新牵头部门、统筹实施主体、运营主体、技术团队及规划、设计、造价、金融、社会、文化、运营等专业配置和职责。",
                     "附组织架构图、责任矩阵、主体授权或会议文件。",
                     "《审查要点》第3、5、15页。"),
                item("1.2.1.2 编制程序与居民参与",
                     "完整记录划定范围、明确主体、片区体检、策划编制、多轮全覆盖意见征集、部门初审、公示、联合评审、政府审定和备案程序。",
                     "形成时间轴、入户/商户/租户台账、公示材料、会议纪要和意见采纳表。",
                     "《审查要点》第2—4、5、13页。"),
            ]),
        ]),
        ("1.3 范围期限", [
            ("1.3.1 片区边界与规模", [
                item("1.3.1.1 边界划定",
                     "综合街道社区边界、自然地理边界、主次干道围合、历史文化资源完整性、10—15分钟生活圈和详细规划单元划定边界；边界不得切割关键资源和实施单元。",
                     "形成区位关系图、片区范围图、边界坐标表及边界合理性说明。",
                     "《审查要点》第2、4、15页。"),
                item("1.3.1.2 规模与期限",
                     "申报片区面积不得低于10公顷；一般城市更新片区按审查表不低于20公顷、县城不低于10公顷，特殊功能片区应说明例外理由。明确现状基期、策划期限和分年度建设期。",
                     "形成面积核算成果、范围审核意见、年度节点表。",
                     "《审查要点》第15页；《试点通知》第2页。"),
            ]),
        ]),
        ("1.4 编制原则", [
            ("1.4.1 原则体系", [
                item("1.4.1.1 六项基本原则",
                     "落实完整连片、规模适度、经济平衡、实施有效、民生优先、保护传承；同步落实运营前置、居民参与、政府投资与社会资本匹配。",
                     "形成原则—章节—项目落实矩阵，避免原则与方案脱节。",
                     "《审查要点》第2、4—6页；《试点通知》第1—3页。"),
                item("1.4.1.2 底线原则",
                     "明确不新增隐性债务、不大拆大建、不破坏历史文脉、不违规征拆、不搞形象工程、不将违法违规行为合法化。",
                     "形成底线清单及项目逐项核验结论。",
                     "《审查要点》第5—8、14页；《试点通知》第2—4页。"),
            ]),
        ]),
        ("1.5 编制依据", [
            ("1.5.1 依据体系与效力", [
                item("1.5.1.1 政策、规划与技术依据",
                     "按法律法规、政策文件、国土空间规划/详细规划、城市更新专项规划、行业专项规划、标准导则和基础数据分组列明，标注发文单位、文号、日期和有效状态。",
                     "形成依据清单及规划套合核验表；拟调详细规划须形成充分论证和部门意见。",
                     "《审查要点》第2、14页。"),
            ]),
        ]),
    ]),
    ("第二章 片区体检", [
        ("2.1 现状概况", [
            ("2.1.1 基础底图与现状画像", [
                item("2.1.1.1 基础数据建库",
                     "统一坐标、时间基准和统计口径，汇集人口、建筑、用地、产权、设施、交通、管网、产业、文化、生态、投资等数据；注明来源、年份、精度和责任单位。",
                     "形成基础数据汇总表、GIS数据库、数据字典和质量核验记录。",
                     "《审查要点》第4、8、12页。"),
                item("2.1.1.2 片区类型判定",
                     "根据主体功能、重点资源、问题结构和发展阶段判定更新类型，并说明与城市更新专项规划传导关系及差异化目标。",
                     "形成类型判定表、现状综合评价图。",
                     "《审查要点》第15页第3项。"),
            ]),
        ]),
        ("2.2 问题调查", [
            ("2.2.1 全要素问题诊断", [
                item("2.2.1.1 民生与空间问题",
                     "调查住房、公共服务、市政基础设施、交通停车、公共空间、无障碍、环境品质等短板，量化服务半径、缺口规模和受影响人群。",
                     "形成问题清单、问题空间分布图和完整社区达标评价。",
                     "《审查要点》第4、7—8、15—16页。"),
                item("2.2.1.2 安全韧性与违建问题",
                     "专项排查消防、燃气、内涝、地灾、建筑安全、地下管网及违法建设；逐点明确风险等级、权属、处置方式和时序。",
                     "形成风险台账、违建点位表/图、地下管网现状图和隐患闭环表。",
                     "《审查要点》第5—7、12、14—16页。"),
            ]),
        ]),
        ("2.3 资源调查", [
            ("2.3.1 更新潜力与权属", [
                item("2.3.1.1 闲置低效与公产资源",
                     "调查闲置建筑、低效用地、边角地、零星空地、公产及可复合利用设施，核实面积、用途、权属、限制条件和盘活意愿。",
                     "形成资源一张图、更新资源表、公产点位表和产权核验材料。",
                     "《审查要点》第5、8、12、15页。"),
                item("2.3.1.2 历史文化与生态资源",
                     "普查历史文化街区、历史建筑、传统风貌、工业遗产、古树、水体、山体等资源，明确保护级别、价值和承载边界。",
                     "形成资源价值评估、保护名录、控制范围图；涉及历史文化街区须编制专篇。",
                     "《审查要点》第6、12、16页。"),
            ]),
        ]),
        ("2.4 意愿调查", [
            ("2.4.1 全覆盖全过程参与", [
                item("2.4.1.1 调查对象与方法",
                     "覆盖住户、商户、租户、产权单位、社区组织和重点群体，采用入户、访谈、座谈、问卷及线上补充等方式，不得以发传单或小样本替代。",
                     "形成样本框、覆盖率统计、原始问卷/访谈记录和空间化诉求图。",
                     "《审查要点》第3、5、7—8、15页。"),
                item("2.4.1.2 意愿转译与反馈",
                     "将居民诉求转译为问题、功能、空间和项目需求，说明采纳、部分采纳或不采纳理由；策划期间至少形成多轮反馈闭环。",
                     "形成意见—方案响应表、版本修改记录、公示意见及合法权益保障说明。",
                     "《审查要点》第3、5、8、13页；《试点通知》第3页。"),
            ]),
        ]),
        ("2.5 运营调查", [
            ("2.5.1 市场与运营基础", [
                item("2.5.1.1 业态与消费需求调查",
                     "调查现有业态、客群、消费频次、租金、空置率、营业时段、周边竞品和交通可达性，识别需求缺口与环境承载上限。",
                     "形成业态POI图、商户访谈、消费客群画像、竞品与租金分析。",
                     "《审查要点》第5、15页。"),
                item("2.5.1.2 运营主体与资源调查",
                     "运营主体全过程参与体检和策划，提出可运营空间、服务场景、主理人/品牌资源和初步经营边界；每项更新内容须对应服务需求和运营场景。",
                     "形成运营主体能力证明、招商资源库、场景清单和入驻意向材料。",
                     "《审查要点》第5、8、16页。"),
            ]),
        ]),
    ]),
    ("第三章 功能业态", [
        ("3.1 功能定位", [
            ("3.1.1 差异化定位与目标", [
                item("3.1.1.1 定位论证",
                     "基于片区类型、区位、资源、产业基础、居民需求和市场需求，明确片区在城市及区域中的差异化角色，不得套用同质化口号。",
                     "形成定位推导链、目标指标表和与专项规划的传导关系。",
                     "《审查要点》第15页第9项。"),
            ]),
        ]),
        ("3.2 业态布置", [
            ("3.2.1 业态体系与空间配置", [
                item("3.2.1.1 主导与复合业态",
                     "确定“1个主导业态+多个复合业态”的初步配比，明确服务对象、规模、载体、运营人和导入时序，校核与周边项目的差异化竞争。",
                     "形成业态配比表、功能业态策划图和供需测算。",
                     "《审查要点》第15页第10项。"),
                item("3.2.1.2 活力空间布局",
                     "业态与交通节点、公共空间、历史文化资源、山水资源和可盘活载体匹配，形成特色业态走廊、活力界面或社区服务网络。",
                     "形成分层业态布局图、首层界面图和昼夜运营时段图。",
                     "《审查要点》第16页第11项。"),
            ]),
        ]),
        ("3.3 盘活空间", [
            ("3.3.1 资源盘活方案", [
                item("3.3.1.1 分类盘活路径",
                     "对闲置低效建筑、场地、边角地和公产资源逐项提出保留利用、改造提升、功能转换、租赁托管、合作共建等方式，说明权属可实施性。",
                     "形成资源—场景—项目—运营主体一一对应表及重点载体概念方案。",
                     "《审查要点》第5—6、8、16—17页。"),
            ]),
        ]),
        ("3.4 运营模式", [
            ("3.4.1 可持续运营机制", [
                item("3.4.1.1 全生命周期运营",
                     "明确建设、招商、开业、稳定运营和更新维护阶段的主体、权责、成本、收入和绩效；公共服务与经营项目分别设计运营模式。",
                     "形成运营架构、业务模型、收入成本表、维护更新机制和退出/调整机制。",
                     "《审查要点》第5、8、16页；《试点通知》第3页。"),
                item("3.4.1.2 运营负面清单",
                     "规避业态同质化、过度商业化、脱离在地文化、规模失衡、扰民和环境承载超负荷等风险。",
                     "形成负面清单、风险阈值及调整预案。",
                     "《审查要点》第16页第13项。"),
            ]),
        ]),
        ("3.5 品牌招商", [
            ("3.5.1 招商前置与落位", [
                item("3.5.1.1 品牌及主理人招商",
                     "同步开展主理人招商，核验品牌与定位、空间和客群的适配性；重点项目不得仅提供意向品牌名单，应形成接洽记录或入驻意向。",
                     "形成招商地图、目标品牌库、接洽台账、意向协议和开业时序。",
                     "《审查要点》第5、16页第12项。"),
            ]),
        ]),
    ]),
    ("第四章 城市设计", [
        ("4.1 整体空间布局", [
            ("4.1.1 空间结构与周边联动", [
                item("4.1.1.1 总体空间方案",
                     "衔接周边功能、交通、生态和公共空间，提出结构清晰、可分期实施的整体空间布局；体现存量提质而非增量扩张。",
                     "形成总平面、空间结构图、周边联动分析和分期实施图。",
                     "《审查要点》第16页第14项。"),
            ]),
        ]),
        ("4.2 综合交通体系", [
            ("4.2.1 路网、慢行与停车", [
                item("4.2.1.1 交通微循环",
                     "打通断头路和微循环，优化出入口与区域交通衔接；保护步行连续性，合理组织消防、公交、配送和旅游交通。",
                     "形成道路等级、交通组织、消防通道、出入口和可达性分析图。",
                     "《审查要点》第5、16页第15项。"),
                item("4.2.1.2 慢行与停车",
                     "构建连续慢行和无障碍系统，基于需求测算配置停车，优先盘活存量和共享资源，避免停车设施过量挤占公共空间。",
                     "形成慢行系统图、无障碍图、停车供需平衡表及共享方案。",
                     "《审查要点》第5、16页。"),
            ]),
        ]),
        ("4.3 地下管网设施", [
            ("4.3.1 管网系统更新", [
                item("4.3.1.1 现状普查与系统方案",
                     "全面排查污水、雨水、供水、燃气、强电、弱电等管网，核实管径、材质、权属、年代、容量和隐患；统筹片区内外及不同产权主体。",
                     "形成现状管网图、问题台账、容量校核和综合管网规划图。",
                     "《审查要点》第6、7—8、16页第16项。"),
                item("4.3.1.2 建设协同",
                     "与相关专项规划、道路和公共空间项目同步设计、同步施工，建立开挖统筹计划，避免反复开挖。",
                     "形成项目接口表、道路—管网联合时序图及产权单位会签意见。",
                     "《审查要点》第6、16页。"),
            ]),
        ]),
        ("4.4 小微空间改造", [
            ("4.4.1 公共空间释放与品质提升", [
                item("4.4.1.1 违建拆除与空间疏导",
                     "在合法合规前提下对违建分类处置，拆除后按居民和公共服务需求释放空间；拆除硬质围墙，盘活边角地、零星地和闲置用地。",
                     "形成违建处置表、拆前拆后空间对比、公共空间增量核算。",
                     "《审查要点》第5、7—8、14、16页。"),
                item("4.4.1.2 小微空间设计",
                     "系统提升社区、公园广场、街巷、近山滨水和文化特色空间，明确使用人群、活动、设施、绿化、照明和维护主体。",
                     "形成小微空间一览表、节点设计图和运营维护卡。",
                     "《审查要点》第5、16页第17项。"),
            ]),
        ]),
        ("4.5 公共服务设施", [
            ("4.5.1 完整社区设施补短板", [
                item("4.5.1.1 配置标准与复合利用",
                     "按完整社区标准和片区类型核算学校、养老、托育、快递、食堂、便民、商业、文化、住房保障等需求，优先利用存量和复合空间。",
                     "形成设施供需表、服务半径图、设施项目表和运营主体清单。",
                     "《审查要点》第7—8、16页第18项。"),
            ]),
        ]),
        ("4.6 市政基础设施", [
            ("4.6.1 市政保障与韧性", [
                item("4.6.1.1 市政设施配置",
                     "结合人口和功能规模，科学配置垃圾收集、公共厕所、二次供水、照明、综合配电、消防、环卫等设施，并校核容量和服务范围。",
                     "形成专项校核表、市政设施布局图和近期改造项目清单。",
                     "《审查要点》第7—8、16页第19项。"),
            ]),
        ]),
        ("4.7 空间管控设计", [
            ("4.7.1 更新方式与风貌管控", [
                item("4.7.1.1 更新对象分类",
                     "科学划定保留、修缮、改造、拆除、功能转换和新建对象，逐对象说明依据；降低不合理建筑密度并释放有效公共空间。",
                     "形成更新方式规划图、对象清单和规划指标对比表。",
                     "《审查要点》第6、16页第21项。"),
                item("4.7.1.2 风貌与历史文化保护",
                     "延续传统肌理和街巷尺度，控制新建、改扩建建筑体量、高度、面宽、色彩和第五立面；活化历史建筑，禁止盲目拔高和风格冲突。",
                     "形成历史文化专篇（适用时）、风貌分区图、控制图则和重点界面设计。",
                     "《审查要点》第6—8、12—13、16页第20—21项；《试点通知》第2—3页。"),
            ]),
        ]),
    ]),
    ("第五章 项目谋划", [
        ("5.1 居住环境改善类", [
            ("5.1.1 住房与社区环境项目", [
                item("5.1.1.1 项目生成与设计深度",
                     "针对住房、公共空间、适老适幼、无障碍和环境品质问题生成项目；每个项目明确对象、建设内容、边界、设计方案、投资估算、资金来源和时序。",
                     "形成项目一案一册、节点设计图及问题—项目销项表。",
                     "《审查要点》第11、17页第22—25项。"),
            ]),
        ]),
        ("5.2 安全韧性类项目", [
            ("5.2.1 风险治理项目", [
                item("5.2.1.1 隐患闭环",
                     "围绕建筑、消防、燃气、内涝、地灾、管网、应急避难等问题形成分级治理项目，优先纳入年度计划和相关部门专项规划。",
                     "形成风险等级—治理项目—责任部门—完成期限闭环表。",
                     "《审查要点》第7—8、11、17页第23、25项。"),
            ]),
        ]),
        ("5.3 功能完善类项目", [
            ("5.3.1 设施补短板项目", [
                item("5.3.1.1 公共服务与市政项目",
                     "将设施缺口转化为可建设、可运营项目，明确服务人口、建设规模、选址、资产权属和运营维护安排。",
                     "形成设施项目卡、专项部门衔接意见和纳规证明。",
                     "《审查要点》第7—8、11、17页第22—23项。"),
            ]),
        ]),
        ("5.4 动能提升类项目", [
            ("5.4.1 产业与消费场景项目", [
                item("5.4.1.1 收益项目策划",
                     "以可盘活载体和真实市场需求为基础，形成主导业态、消费场景、招商对象、建设改造和运营方案；收益类项目原则上引入社会资本。",
                     "形成商业模型、招商意向、社会资本参与方案和项目现金流。",
                     "《审查要点》第5—6、11、17页第24、27项；《试点通知》第1、3页。"),
            ]),
        ]),
        ("5.5 生态改善类项目", [
            ("5.5.1 生态修复与绿色基础设施", [
                item("5.5.1.1 山水空间与海绵项目",
                     "统筹山体、水体、绿地、雨洪和热环境，优先采用小尺度、低影响、可维护措施，避免脱离片区实际的大景观工程。",
                     "形成生态问题图、绿色基础设施网络、海绵指标和运维方案。",
                     "《审查要点》第4—6、11、16页第17项。"),
            ]),
        ]),
        ("5.6 文化保护利用类项目", [
            ("5.6.1 保护修缮与活化利用", [
                item("5.6.1.1 文化项目闭环",
                     "坚持保护优先，明确保护对象、修缮标准、活化功能、运营主体和客流承载；不得拆除历史建筑或破坏街区格局。",
                     "形成价值评估、修缮设计、活化运营方案及保护主管部门意见。",
                     "《审查要点》第6—8、11—13、16页；《试点通知》第2—3页。"),
            ]),
        ]),
    ]),
    ("第六章 效益分析", [
        ("6.1 项目成本测算", [
            ("6.1.1 全生命周期成本", [
                item("6.1.1.1 成本口径与估算",
                     "统一估算基准和价格时点，覆盖前期、征拆/权益协调、建安、设备、工程其他费、预备费、融资、招商、运营启动、维护和更新成本；避免漏项和重复。",
                     "形成分项目估算表、测算依据、生命周期成本表和不确定性说明。",
                     "《审查要点》第9、17页第26项。"),
            ]),
        ]),
        ("6.2 资金筹集计划", [
            ("6.2.1 分类匹配资金", [
                item("6.2.1.1 公益安全类资金",
                     "优先对接中央和省级财政资金及相关部门专项资金，保持渠道不乱、用途不变；明确申报条件、责任单位、到位时点和缺口。",
                     "形成资金渠道清单、申报计划、部门确认和年度到位表。",
                     "《审查要点》第6、8—9、17页第27项；《试点通知》第2、4—5页。"),
                item("6.2.1.2 收益类资金",
                     "以社会资本和项目经营收益为主，明确投资主体、资本金、融资边界、回收机制和风险分担；融资项目应充分论证，必要时取得省级金融机构书面认可。",
                     "形成投融资结构图、社会资本意向、金融机构意见和合规审查记录。",
                     "《审查要点》第6、8—9、17页第24、27项。"),
            ]),
        ]),
        ("6.3 资金平衡分析", [
            ("6.3.1 片区与项目平衡", [
                item("6.3.1.1 现金流与平衡边界",
                     "分别测算公益安全类与收益类项目，禁止以不确定土地收益、过度融资或债务转嫁实现账面平衡；开展建设成本、租金、出租率和资金到位时序敏感性分析。",
                     "形成分年度现金流、资金缺口、敏感性分析、压力测试和财政承受说明。",
                     "《审查要点》第5—9、14、17页；《试点通知》第2—3、5页。"),
                item("6.3.1.2 多元参与机制",
                     "审慎论证产权置换、租赁托管、合作共建、以资入股、委托运营、收益分成，以及居民直接出资、让渡权益、住房公积金等路径，确保自愿、合法、可执行。",
                     "形成权责收益表、协议框架、产权人意愿证明和法律合规意见。",
                     "《审查要点》第6、17页第26项。"),
            ]),
        ]),
        ("6.4 社会效益分析", [
            ("6.4.1 综合绩效与示范推广", [
                item("6.4.1.1 绩效指标",
                     "量化民生改善、安全韧性、公共空间、设施补齐、就业消费、文化保护、生态环境、居民满意度、社会资本撬动及资产盘活成效，设置基期、目标值、监测频次和责任单位。",
                     "形成绩效指标表、监测数据来源和2027年中期/期末目标。",
                     "《试点通知》第1、4页。"),
                item("6.4.1.2 可复制模式",
                     "提炼机制、资金、运营、治理和技术路径的适用条件、成本边界和复制步骤，不得仅作宣传性总结。",
                     "形成模式卡、标准流程和推广清单。",
                     "《试点通知》第1、4页。"),
            ]),
        ]),
    ]),
    ("第七章 实施保障", [
        ("7.1 组织架构", [
            ("7.1.1 专班与主体责任", [
                item("7.1.1.1 组织责任矩阵",
                     "建立政府指导、住建牵头、部门联动、属地落实、统筹实施主体组织、运营主体全过程参与的机制，明确决策、审批、建设、资金、招商、运营和监管责任。",
                     "形成组织架构图、RACI责任矩阵、议事规则和部门联络表。",
                     "《审查要点》第3、5—7、15页；《试点通知》第3页。"),
            ]),
        ]),
        ("7.2 建设时序", [
            ("7.2.1 分年度实施计划", [
                item("7.2.1.1 项目排期与资金协同",
                     "制定分年度项目计划，优先安排民生公益安全类项目；项目成熟度、审批、资金到位、施工条件和运营导入必须协同。",
                     "形成年度甘特图、前置条件清单、资金到位计划和开工准备度评价。",
                     "《审查要点》第17页第25项。"),
                item("7.2.1.2 试点节点控制",
                     "获支持片区在补助资金到位后3个月内启动建设；为2027年8月中期评估和2027年底明显成效倒排策划、项目库、实施方案和开工节点。",
                     "形成里程碑计划、月度责任清单和节点预警机制。",
                     "《试点通知》第4页。"),
            ]),
        ]),
        ("7.3 工作机制", [
            ("7.3.1 审批、协调与动态管理", [
                item("7.3.1.1 政策与审批协同",
                     "梳理用地、规划、建设、消防、历史文化、资金等审批事项和地方支持政策，建立跨部门问题清单、会商机制和限时销项制度。",
                     "形成审批路线图、政策工具箱、会商纪要和销项台账。",
                     "《审查要点》第3—4、6页；《试点通知》第3、5页。"),
                item("7.3.1.2 项目库动态管理",
                     "建立谋划库、储备库、在建库和运营库，统一项目编码、状态、投资、资金、责任和绩效字段，按成熟度动态进退库。",
                     "形成项目库数据表、入库标准、月度更新机制和省级项目库报送接口。",
                     "《审查要点》第4页；《试点通知》第4—5页。"),
            ]),
        ]),
        ("7.4 长效运营", [
            ("7.4.1 建管运一体化", [
                item("7.4.1.1 运营交付与绩效",
                     "在设计和建设阶段嵌入运营条件，明确资产移交、招商开业、设施维护、公共服务购买、运营补贴边界和绩效考核。",
                     "形成建管运接口表、资产移交清单、运营KPI和年度预算。",
                     "《审查要点》第3、5、8、12页；《试点通知》第3页。"),
            ]),
        ]),
        ("7.5 社会治理", [
            ("7.5.1 居民共治与风险管理", [
                item("7.5.1.1 共同参与机制",
                     "建立市场主体、社会资本、产权人（居民）、社区组织共同参与机制，明确意见表达、权益协商、争议处理和运营监督渠道。",
                     "形成居民议事规则、权益保障方案、投诉处置流程和满意度回访机制。",
                     "《审查要点》第3、5—6页；《试点通知》第3页。"),
                item("7.5.1.2 监测评估与风险预警",
                     "围绕建设任务、政策落实、预算执行、资金使用、机制创新、运营成效和群众权益建立动态监测；对隐债、征拆、文化破坏、形象工程、资金滞留等风险设置预警。",
                     "形成监测指标表、季度评估、风险台账、纠偏和问责机制。",
                     "《试点通知》第3—4页。"),
            ]),
        ]),
    ]),
]


def add_requirement_block(doc: Document, requirement: str, evidence: str, policy: str) -> None:
    p = doc.add_paragraph(style="Framework Requirement")
    r = p.add_run("工作要求｜")
    set_run_font(r, size=9.5, color=BLUE, bold=True)
    r = p.add_run(requirement)
    set_run_font(r, size=9.5, color=TEXT)

    p = doc.add_paragraph(style="Evidence Line")
    r = p.add_run("最低成果证据｜")
    set_run_font(r, size=9, color=GREEN, bold=True)
    r = p.add_run(evidence)
    set_run_font(r, size=9, color=GREEN)

    p = doc.add_paragraph(style="Policy Line")
    r = p.add_run("政策依据｜")
    set_run_font(r, size=8.5, color=MUTED, bold=True)
    r = p.add_run(policy)
    set_run_font(r, size=8.5, color=MUTED)


def add_framework(doc: Document) -> None:
    for chapter_title, sections in CHAPTERS:
        doc.add_heading(chapter_title, level=1)
        for section_title, subsections in sections:
            doc.add_heading(section_title, level=2)
            for subsection_title, items in subsections:
                doc.add_heading(subsection_title, level=3)
                for title, requirement, evidence, policy in items:
                    doc.add_heading(title, level=4)
                    add_requirement_block(doc, requirement, evidence, policy)


def add_appendix_tables(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("附录A 报审前一票否决核验表", level=1)
    p = doc.add_paragraph("以下任一项结论为“否”时，不进入正式报审程序；项目负责人应组织整改并留存销项证据。")
    redlines = [
        ("A01", "不新增地方政府隐性债务", "财政/审计合规意见；资金方案"),
        ("A02", "不随意扩大融资规模", "融资边界和偿债来源论证"),
        ("A03", "片区及分项目能够实现资金平衡", "现金流、敏感性与压力测试"),
        ("A04", "不存在大拆大建和破坏历史文脉", "更新方式图；文化保护专项"),
        ("A05", "已实质性开展片区体检，不以城市体检替代", "片区体检报告及原始数据库"),
        ("A06", "已明确统筹实施主体", "政府或主管部门文件"),
        ("A07", "已落实运营主体并全过程参与", "主体证明、会议/招商记录"),
        ("A08", "已组建工作专班并建立部门协同机制", "专班文件、责任矩阵"),
        ("A09", "已全面调查违建并按需释放公共空间", "违建台账、处置方案"),
        ("A10", "详细规划调整建议已充分论证并征求部门意见", "论证报告、部门书面意见"),
        ("A11", "文本、附图、附表和附件齐全完整", "完整性清单"),
        ("A12", "不存在违规征拆、侵害群众权益、形象工程等重大风险", "权益保障和风险审查意见"),
    ]
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["编号", "核验事项", "最低证据", "结论", "责任人/日期"]
    widths = [Cm(1.2), Cm(6.1), Cm(4.7), Cm(1.6), Cm(2.6)]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=WHITE, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[i], RED)
    set_repeat_table_header(table.rows[0])
    for idx, (code, check, evidence) in enumerate(redlines):
        row = table.add_row()
        values = (code, check, evidence, "□是  □否", "")
        for i, value in enumerate(values):
            set_cell_text(row.cells[i], value, bold=(i == 0), size=8.2)
            if idx % 2:
                set_cell_shading(row.cells[i], "FFF7F7")
        prevent_row_split(row)
    set_table_borders(table, color="D9A6A6")

    doc.add_heading("附录B 方案审查28项对照表", level=1)
    review_rows = [
        (1, "片区规模", "1.3、1.4", "面积、城市能级与示范性"),
        (2, "片区边界", "1.3", "行政/自然/道路/文化/生活圈/详规单元"),
        (3, "类型判定", "2.1、3.1", "主体功能、资源特点、差异化目标"),
        (4, "组织架构", "1.2、7.1", "专班、牵头、统筹、运营主体权责"),
        (5, "问题调查", "2.2", "地下管网、违建等问题全面准确"),
        (6, "资源调查", "2.3", "闲置低效、历史文化、公产权属"),
        (7, "意愿调查", "2.4", "居民全覆盖、全流程参与"),
        (8, "运营调查", "2.5", "业态、功能适配、消费需求"),
        (9, "功能定位", "3.1", "区位资源、产业基础、市场与专项规划"),
        (10, "业态配置", "3.2", "1主导+多复合、初步配比、差异化"),
        (11, "空间布局", "3.2、3.3", "交通/公共空间/资源匹配、活力界面"),
        (12, "运营招商", "3.4、3.5", "主体能力、品牌匹配、主理人意向"),
        (13, "负面清单", "3.4", "同质化、过度商业化、文化与承载风险"),
        (14, "整体空间布局", "4.1", "周边联动、空间优化、风貌协同"),
        (15, "综合交通体系", "4.2", "路网织补、慢行、停车、区域衔接"),
        (16, "地下管网设施", "4.3", "全面排查、专项衔接、系统治理"),
        (17, "小微空间改造", "4.4", "社区/公园/街巷/山水/文化空间"),
        (18, "公共服务设施", "4.5", "完整社区标准与复合配置"),
        (19, "市政基础设施", "4.6", "环卫、供水、照明、配电等容量"),
        (20, "历史文化保护", "2.3、4.7、5.6", "资源挖掘、活化、风貌控制"),
        (21, "空间管控设计", "4.7", "保留/改造/拆除/转换及公共空间"),
        (22, "项目内容", "第五章", "建设内容、估算、时序、资金来源"),
        (23, "项目衔接", "第二、五章", "体检更新一体、部门项目整合"),
        (24, "项目深度", "第五章", "设计方案深度、社会资本参与"),
        (25, "建设时序", "5、7.2", "分年度项目、公益安全优先、资金协同"),
        (26, "成本测算", "6.1、6.3", "全生命周期、多元产权与居民参与"),
        (27, "资金筹集", "6.2、6.3", "公益财政、收益社会资本、融资认可"),
        (28, "其他", "全篇", "结合城市和片区实际补充"),
    ]
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["序号", "审查要点", "对应章节", "内部核验重点", "状态"]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=WHITE, size=8.4, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[i], NAVY)
    set_repeat_table_header(table.rows[0])
    for idx, values in enumerate(review_rows):
        row = table.add_row()
        for i, value in enumerate((*values, "□")):
            set_cell_text(row.cells[i], str(value), bold=(i == 0), size=8.0)
            if idx % 2:
                set_cell_shading(row.cells[i], "F7F9FB")
        prevent_row_split(row)
    set_table_borders(table)

    doc.add_heading("附录C 必备附图、附表及附件清单", level=1)
    deliverables = [
        ("附图", "区位关系图；片区范围图；上位规划衔接图；问题空间分布图；违建点位分布图；更新潜力分析图；功能业态策划图；更新方式规划图；更新片区设计图；更新项目分布图；其他专项设计图。"),
        ("附表", "片区基础数据汇总表；片区体检问题汇总表；片区违建点位汇总表；片区公产点位汇总表；片区更新资源汇总表；片区更新项目汇总表；其他汇总表。"),
        ("项目表最低字段", "项目名称、项目编码、位置/边界、问题来源、建设内容、设计方案、投资估算、资金来源、建设时序、责任主体、运营主体、审批状态、绩效指标。"),
        ("附件", "专项规划设计方案；征求意见情况；策划方案公示情况；相关会议纪要；评审意见和修改落实情况；其他有关材料。"),
        ("条件性专篇", "片区含历史文化街区时编制历史文化专篇；采用自主更新模式时编制自主更新专篇；有条件的片区附完整运营方案。"),
        ("公示材料", "更新范围及规模、更新方式、更新定位、核心控制指标、方案总平面及局部效果图；附官网/公众号截图、社区张贴照片、群众书面建议等。"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(("类别", "明确工作要求")):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=WHITE, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[i], BLUE)
    set_repeat_table_header(table.rows[0])
    for idx, values in enumerate(deliverables):
        row = table.add_row()
        set_cell_text(row.cells[0], values[0], bold=True, size=8.5)
        set_cell_text(row.cells[1], values[1], size=8.5)
        if idx % 2:
            set_cell_shading(row.cells[0], "F7F9FB")
            set_cell_shading(row.cells[1], "F7F9FB")
        prevent_row_split(row)
    set_table_borders(table)

    doc.add_heading("附录D 2027年省级财政支持申报材料清单", level=1)
    support = [
        ("D01", "片区策划方案", "按本框架完成；通过完整性、合规性和28项审查内审"),
        ("D02", "城市更新工作材料", "年度工作要点、工作方案"),
        ("D03", "城市更新专项规划", "正式规划或“十五五”规划过程稿，并说明技术团队、编制情况和计划"),
        ("D04", "城市体检材料", "城市体检报告和独立的片区体检报告"),
        ("D05", "项目库材料", "城市更新谋划库、储备库、在建库等"),
        ("D06", "地方支持政策", "用地、审批、资金等现行地方政策文件"),
        ("D07", "资金保障及财政支撑", "2024—2026年中央、省级财政资金申报及使用情况；不新增隐债承诺"),
        ("D08", "报送与保密", "纸质材料3份；同步电子版；合计原则上不超过4GB；不得含涉密内容"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(("编号", "材料", "完成标准", "责任/状态")):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=WHITE, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[i], NAVY)
    set_repeat_table_header(table.rows[0])
    for idx, values in enumerate(support):
        row = table.add_row()
        for i, value in enumerate((*values, "")):
            set_cell_text(row.cells[i], value, bold=(i == 0), size=8.2)
            if idx % 2:
                set_cell_shading(row.cells[i], "F7F9FB")
        prevent_row_split(row)
    set_table_borders(table)

    doc.add_heading("附录E 编制成果提交前检查签署页", level=1)
    checks = [
        ("技术负责人", "四级框架内容完整；图、表、文、项目库数据一致", ""),
        ("片区体检负责人", "调查真实、全覆盖、可追溯；问题与项目闭环", ""),
        ("运营负责人", "运营主体落实；场景、招商、现金流和长效运营可行", ""),
        ("造价/财务负责人", "估算口径一致；资金来源可靠；资金平衡及压力测试通过", ""),
        ("规划负责人", "符合上位规划；详细规划调整建议论证充分", ""),
        ("历史文化负责人", "保护底线核验通过；无大拆大建和文脉破坏", ""),
        ("统筹实施主体", "主体、专班、居民意见、部门意见和实施条件落实", ""),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(("签署角色", "签署结论", "签字", "日期")):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=WHITE, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[i], BLUE)
    set_repeat_table_header(table.rows[0])
    for idx, (role, conclusion, _) in enumerate(checks):
        row = table.add_row()
        for i, value in enumerate((role, conclusion, "", "")):
            set_cell_text(row.cells[i], value, bold=(i == 0), size=8.3)
            if idx % 2:
                set_cell_shading(row.cells[i], "F7F9FB")
        prevent_row_split(row)
    set_table_borders(table)

    doc.add_heading("附录F 片区范围未确定阶段工作包", level=1)
    p = doc.add_paragraph(
        "遵义基础资料目录现有约2766个文件，包含城市体检成果、指标体系、体检单元、"
        "城市更新专项规划、既有片区策划、项目库、GIS数据库、地下管网和历史文化资料。"
        "范围未确定不等于项目停工，应采用“全市底盘先行—候选片区筛选—边界方案比选—确认后深化”的路径。"
    )
    work_packages = [
        ("F01", "资料目录与证据台账", "立即开展", "全量文件索引、去重、版本/年份/空间范围识别、证据A—E分级；建立原件只读区和清洗区。", "资料总目录、数据字典、缺失/冲突清单"),
        ("F02", "政策与指标字典", "立即开展", "整合审查要点、城市体检手册、贵州指标体系、更新专项规划和地方政策，拆解指标定义、分子分母、数据源与责任部门。", "政策矩阵、指标字典、部门资料函"),
        ("F03", "全市空间底盘", "立即开展", "整理行政区、街道社区、体检单元、道路、建筑、设施、项目、历史文化、地下管网等图层，统一坐标和主键。", "GIS图层目录、空间主键表、基础底图"),
        ("F04", "既有体检成果复核", "立即开展", "将2025基线、2026成果和社区补充调查按年份分层，识别复制值、口径冲突和待补数项；历史数据不得直接改写为当年结论。", "指标状态表、红色校核项、可引用/线索/待核清单"),
        ("F05", "候选片区初筛", "立即开展", "利用体检问题、更新潜力、项目成熟度、历史文化、空间连片性和运营资源，对既有更新单元及老火车站、老城—丁字口、国贸商圈、插旗山等资料覆盖区域进行桌面初筛。", "候选片区长名单、初筛评分表、问题与资源热力图"),
        ("F06", "既有项目库空间化", "立即开展", "整理市级项目库、成熟项目库、专项资金项目和既有策划项目，统一项目编码、边界、投资、资金、进度和责任主体。", "项目一张图、项目成熟度分级、重复/缺项清单"),
        ("F07", "运营与资金资源预研", "立即开展", "梳理可运营存量载体、政府/国企资源、潜在运营主体、消费场景及中央省级资金适配条件，只形成资源线索，不先行承诺收益。", "运营资源库、资金渠道矩阵、访谈对象清单"),
        ("F08", "调查与调研工具准备", "立即开展", "设计住户、商户、租户、产权单位和运营主体问卷；制定照片编码、点位采集、管网核验和安全隐患记录标准。", "问卷与访谈提纲、外业手册、调研底图模板"),
        ("F09", "候选片区踏勘", "形成长名单后", "对候选区域开展快速踏勘和样本访谈，核实连片性、主要问题、资源、实施主体和运营条件；不得替代最终全覆盖调查。", "踏勘记录、样本照片、候选区事实卡"),
        ("F10", "边界方案比选", "形成长名单后", "每个重点候选区域形成不少于2—3个边界方案，比较面积、边界完整性、生活圈、问题集中度、资源完整性、项目成熟度和资金平衡潜力。", "边界方案图、比选矩阵、推荐及风险说明"),
        ("F11", "片区体检正式实施", "范围确认后", "锁定统计时点和对象清单，开展全覆盖问题、资源、意愿和运营调查，计算正式指标并完成四证合一问题认定。", "正式片区体检报告、数据库和问题清单"),
        ("F12", "策划与项目深化", "范围确认后", "将经审核问题转译为功能、空间、项目、资金和运营方案，完成设计、估算、年度时序及资金平衡。", "正式策划方案、项目一案一册和资金平衡表"),
    ]
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["编号", "工作包", "启动条件", "明确工作要求", "阶段成果"]
    widths = [Cm(1.2), Cm(2.6), Cm(2.2), Cm(6.2), Cm(3.8)]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].width = widths[i]
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=WHITE, size=8.3, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[i], NAVY)
    set_repeat_table_header(table.rows[0])
    for idx, values in enumerate(work_packages):
        row = table.add_row()
        for i, value in enumerate(values):
            row.cells[i].width = widths[i]
            set_cell_text(row.cells[i], value, bold=(i in (0, 1)), size=7.8)
            if idx % 2:
                set_cell_shading(row.cells[i], "F7F9FB")
        prevent_row_split(row)
    set_table_borders(table)

    doc.add_heading("范围确定前不得定稿的事项", level=2)
    blocked = [
        "不得确定正式片区统计分母、指标结果及达标/不达标结论。",
        "不得以候选区样本访谈代替住户、商户、租户等多轮全覆盖意愿调查。",
        "不得形成最终问题清单、项目边界、建设规模、投资估算和资金平衡结论。",
        "不得确定最终功能业态配比、详细城市设计方案和主理人入驻承诺。",
        "不得启动正式公示、联合评审和政府审定程序。",
    ]
    for text in blocked:
        p = doc.add_paragraph(style="Framework Requirement")
        r = p.add_run("控制要求｜")
        set_run_font(r, size=9.5, color=RED, bold=True)
        r = p.add_run(text)
        set_run_font(r, size=9.5, color=TEXT)

    doc.add_paragraph()
    p = doc.add_paragraph(style="Policy Line")
    r = p.add_run("资料来源说明｜")
    set_run_font(r, size=8.5, color=MUTED, bold=True)
    r = p.add_run(
        "本框架依据用户提供的《贵州省城市更新片区策划方案审查要点（试行）》"
        "（贵州省住房和城乡建设厅，2026年6月26日）及2027年省级财政资金支持城市更新储备项目遴选通知"
        "（贵州省住房和城乡建设厅，2026年7月22日）编制，并引用“2026-716-New project”中的片区体检任务书解析、"
        "初始化控制及社区维度成果补充稿作为方法支持。正式报审时应同步核验项目所在地最新有效政策和经确认的指标体系。"
    )
    set_run_font(r, size=8.5, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    section.orientation = WD_ORIENT.PORTRAIT
    configure_styles(doc)
    add_header_footer(doc)
    core = doc.core_properties
    core.title = "贵州省城市更新片区策划报告工作框架"
    core.subject = "四级目录、六级标题样式、政策审查闭环"
    core.author = "Codex"
    core.keywords = "贵州省, 城市更新, 片区策划, 报告框架, 审查要点"


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_title_page(doc)
    add_summary_tables(doc)
    add_framework(doc)
    add_appendix_tables(doc)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH.resolve())


if __name__ == "__main__":
    main()
